# -*- coding: utf-8 -*-
"""Build the fill-in upload templates from DATA_SCHEMAS.

One-off generator (the retired `.workbuddy\\make_upload_templates.py`). The
schema is the single definition of each upload file's columns, so the templates
are derived from it - never typed twice.

Writes:
    Upload_Template_MaterialCodes.xlsx
    Upload_Template_ShipToLocations.xlsx     (no Domain column any more)
    Upload_Template_Domains.xlsx             (new: the domain table)

Run:  python tools/make_upload_templates.py
"""
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                "..", "src"))
import orderflow  # noqa: E402

BASE = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                    ".."))

# schema key -> (file name, sheet name)
TARGETS = {
    "materials": ("Upload_Template_MaterialCodes.xlsx", "Material Codes"),
    "customers": ("Upload_Template_ShipToLocations.xlsx", "Ship To Locations"),
    "domains":   ("Upload_Template_Domains.xlsx", "Domains"),
}

BLANK_ROWS = 200          # the empty rows the user types into


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    de = orderflow.data_export
    for name, (fname, sheet) in TARGETS.items():
        spec = de.schema(name)
        cols = spec["columns"]

        bad = de.contract_ok(name)
        if bad:
            print(f"  {name}: contract check failed, not written")
            for b in bad:
                print("     ", b)
            return 1

        wb = Workbook()
        ws = wb.active
        ws.title = sheet

        head_fill = PatternFill("solid", fgColor="DDEBF7")
        head_font = Font(bold=True, color="1F3864")
        line = Side(style="thin", color="B4C6E7")
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(1, j, c["header"])
            cell.fill, cell.font = head_fill, head_font
            cell.border = Border(bottom=line)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = c.get("width", 18)
            if c.get("title"):
                from openpyxl.comments import Comment
                cell.comment = Comment(c["title"], "OrderFlow")
            # Every data cell is text except a column the importer reads as a
            # number: a leading zero must survive, and a long account number
            # must not lose its last digits to a float.
            fmt = "General" if c.get("kind") == de.KIND_NUMBER else "@"
            for r in range(2, BLANK_ROWS + 2):
                ws.cell(r, j).number_format = fmt
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "A2"

        out = os.path.join(BASE, fname)
        wb.save(out)
        print(f"  wrote {fname}: {len(cols)} column(s) -> "
              f"{', '.join(c['header'] for c in cols)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
